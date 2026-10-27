"""Database-style Excel report for a focused single-run assessment."""

from pathlib import Path

from data.run import MonteCarloRun
from single_run.damping_stability import TransferDampingAnalysis
from single_run.initialization import InitializationAnalysis
from single_run.high_frequency import HighFrequencyAnalysis
from single_run.voltage_stability import VoltageStabilityAnalysis
from single_run.plant_recovery import PlantRecoveryAnalysis
from stability.angular import ANGULAR_REFERENCE_NOTE, AngularStabilityAnalysis
from stability.dynamic_stability import DynamicStabilityAssessment
from stability.scenario import EvaluationProfile


class SingleRunReport:
    """Write concise conclusions plus traceable, relational-style details."""

    def __init__(self, run: MonteCarloRun, profile: EvaluationProfile, initialization: InitializationAnalysis, voltage: VoltageStabilityAnalysis, damping: TransferDampingAnalysis, high_frequency: HighFrequencyAnalysis, angular: AngularStabilityAnalysis, plant_recovery: PlantRecoveryAnalysis, spatial_stability, dynamic: DynamicStabilityAssessment) -> None:
        self.run = run
        self.profile = profile
        self.initialization = initialization
        self.voltage = voltage
        self.damping = damping
        self.high_frequency = high_frequency
        self.angular = angular
        self.plant_recovery = plant_recovery
        self.spatial_stability = spatial_stability
        self.dynamic = dynamic

    def write(self, output_directory: str | Path) -> Path:
        """Create the workbook at the root of the run analysis outputs."""
        from openpyxl import Workbook

        output_path = Path(output_directory) / "analysis_report.xlsx"
        workbook = Workbook()
        sheets = {
            "Summary": workbook.active,
            "Initialization": workbook.create_sheet("Initialization"),
            "Voltage Assessment": workbook.create_sheet("Voltage Assessment"),
            "Transfer Assessment": workbook.create_sheet("Transfer Assessment"),
            "EM Modes": workbook.create_sheet("EM Modes"),
            "Damping Pairs": workbook.create_sheet("Damping Pairs"),
            "High-Frequency Components": workbook.create_sheet("High-Frequency Components"),
            "Angular Assessment": workbook.create_sheet("Angular Assessment"),
            "Plant Recovery": workbook.create_sheet("Plant Recovery"),
            "Spatial Stability": workbook.create_sheet("Spatial Stability"),
            "Definitions": workbook.create_sheet("Definitions"),
        }
        sheets["Summary"].title = "Summary"
        self._write_summary(sheets["Summary"])
        self._write_initialization(sheets["Initialization"])
        self._write_voltage(sheets["Voltage Assessment"])
        self._write_transfers(sheets["Transfer Assessment"])
        self._write_modes(sheets["EM Modes"])
        self._write_pairs(sheets["Damping Pairs"])
        self._write_high_frequency(sheets["High-Frequency Components"])
        self._write_angular(sheets["Angular Assessment"])
        self._write_plant_recovery(sheets["Plant Recovery"])
        self._write_spatial_stability(sheets["Spatial Stability"])
        self._write_definitions(sheets["Definitions"])
        for sheet in workbook.worksheets:
            self._format_sheet(sheet)
        workbook.save(output_path)
        workbook.close()
        return output_path

    def _write_summary(self, sheet) -> None:
        voltage_total = len(self.voltage.assessments)
        voltage_stable = sum(item.assessment.is_stable for item in self.voltage.assessments)
        initialized_count = sum(item.assessment.is_well_initialized for item in self.initialization.records)
        statuses = tuple(record.assessment.status for record in self.damping.records)
        high_frequency_warnings = len(self.high_frequency.warning_records)
        failed_buses = tuple(
            item.bus.name
            for item in self.voltage.assessments
            if not item.assessment.is_stable
        )
        limiting_transfers = tuple(
            record.assessment.transfer_name
            for record in self.damping.records
            if record.assessment.status in {"damped_under", "growing"}
        )
        angular_violations = tuple(
            machine.machine_name
            for machine in self.angular.machines
            if not machine.is_within_limit
        )
        plant_recovery_exceptions = self.plant_recovery.non_recovered_records
        plant_recovery_names = tuple(
            f"{record.plant_name} ({record.assessment.status})"
            for record in plant_recovery_exceptions
        )
        sheet.append(("Scope", "Item", "Value", "Detail"))
        rows = (
            ("Case", "Result directory", self.run.display_name, None),
            ("Case", "Scenario", self.profile.scenario.name, None),
            ("Case", "Severity", self.profile.scenario.severity, None),
            ("Case", "Fault / clearance [s]", f"{self.profile.scenario.fault_time_s:.3f} / {self.profile.scenario.clearance_time_s:.3f}", None),
            ("Dynamic Stability", "Assessment", self.dynamic.status, self.dynamic.reason),
            ("Dynamic Stability", "Voltage criterion", self.dynamic.voltage_status, "All assessed 220/500 kV buses must comply."),
            ("Dynamic Stability", "EM damping criterion", self.dynamic.modal_damping_status, "See Transfer Assessment for the mode-level trace."),
            ("Dynamic Stability", "Angular criterion", self.dynamic.angular_status, "COI-referenced machine excursions."),
            ("Initialization", "Assessed 220/500 kV voltage signals", len(self.initialization.records), None),
            ("Initialization", "Well-initialized voltage signals", initialized_count, None),
            ("Initialization", "Initialization warnings", len(self.initialization.records) - initialized_count, None),
            ("Voltage", "220/500 kV assessed buses", voltage_total, None),
            ("Voltage", "Buses passing all active criteria", voltage_stable, None),
            ("Voltage", "Buses failing one or more criteria", voltage_total - voltage_stable, None),
            ("Voltage", "Failing buses", ", ".join(failed_buses) or "None", "Detailed criteria are in Voltage Assessment."),
            ("Damping", "Published active-power transfers", len(statuses), None),
            ("Damping", "Damped", statuses.count("damped"), None),
            ("Damping", "Damped under 5%", statuses.count("damped_under"), None),
            ("Damping", "Growing", statuses.count("growing"), None),
            ("Damping", "Not measurable", statuses.count("not_measurable"), None),
            ("Damping", "Not applicable", statuses.count("opened") + statuses.count("inactive_pre_fault"), None),
            ("Damping", "Transfers limiting EM damping", ", ".join(limiting_transfers) or "None", "A limiting mode has damping below 5% or is growing."),
            ("High frequency", "Transfer warnings", high_frequency_warnings, "Observation only; does not alter dynamic stability."),
            ("Angular", "COI-referenced machine excursions", len(self.angular.machines), ANGULAR_REFERENCE_NOTE),
            ("Angular", "Angular stability", "Stable" if self.angular.is_stable else "Unstable", "All machines must remain strictly within ±120°"),
            ("Angular", "Machines within ±120°", sum(machine.is_within_limit for machine in self.angular.machines), None),
            ("Angular", "Violating machines", ", ".join(angular_violations) or "None", "Detailed excursions are in Angular Assessment."),
            ("Plant recovery", "Assessed PV/WF/BESS plants", len(self.plant_recovery.records), "Complete P/Q signal pairs only."),
            ("Plant recovery", "Recovered at tmax", len(self.plant_recovery.records) - len(plant_recovery_exceptions), "P and Q both within their terminal tolerance."),
            ("Plant recovery", "Terminal recovery exceptions", len(plant_recovery_exceptions), "Does not alter the dynamic-stability assessment."),
            ("Plant recovery", "Plants outside terminal tolerance", ", ".join(plant_recovery_names) or "None", "Detailed P/Q values are in Plant Recovery."),
            ("Spatial stability", "Map generation", self.spatial_stability.status, self.spatial_stability.reason),
            ("Spatial stability", "Generated maps", len(self.spatial_stability.figure_paths), "Spatial metadata and any skip reason are in 06_Spatial_Stability."),
        )
        for row in rows:
            sheet.append(row)

    def _write_initialization(self, sheet) -> None:
        sheet.append(("Bus", "Voltage level [kV]", "Reference voltage [pu]", "P99 fast error [%]", "Fast RMS error [%]", "Assessment", "Reason"))
        for item in self.initialization.records:
            value = item.assessment
            sheet.append((item.bus.name, item.voltage_level_kv, value.reference_value, value.fast_relative_error_percentile_99 * 100, value.fast_rms_relative_error * 100, "Well initialized" if value.is_well_initialized else "Warning", ", ".join(value.rejection_reasons) or None))

    def _write_voltage(self, sheet) -> None:
        sheet.append(("Bus", "Voltage level [kV]", "Assessment", "0.70 pu compliant", "0.70 pu event", "0.70 pu decision time [s]", "0.70 pu duration [s]", "Below 0.80 pu duration [s]", "0.80 pu compliant", "Settlement compliant", "Final peak error [%]", "Final range [%]", "Fault-on area [pu*s]", "Recovery area [pu*s]", "Unified response area [pu*s]"))
        for item in self.voltage.assessments:
            value = item.assessment
            sheet.append((item.bus.name, item.voltage_level_kv, "Stable" if value.is_stable else "Unstable", value.is_crisis_voltage_compliant, value.crisis_decision_event, value.crisis_decision_time_s, value.crisis_decision_duration_s, value.cumulative_low_voltage_duration_s, value.is_low_voltage_duration_compliant, value.is_settled, value.final_settlement_peak_absolute_relative_error * 100, value.final_settlement_peak_to_peak_relative_range * 100, value.fault_on_area_pu_s, value.recovery_area_pu_s, value.unified_response_area_pu_s))

    def _write_transfers(self, sheet) -> None:
        sheet.append((
            "Transfer",
            "Pre-fault P [MW]",
            "Final P [MW]",
            "Transfer assessment",
            "Detected EM modes",
            "Measurable EM modes",
            "EM composition",
            "Representative mode [Hz]",
            "Representative damping [%]",
            "Representative mode status",
            "Limiting mode [Hz]",
            "Limiting damping [%]",
            "Limiting mode status",
            "Damping pairs",
            "Decision basis",
            "Note",
        ))
        for record in self.damping.records:
            value = record.assessment
            dominant, critical = value.dominant_mode, value.critical_mode
            measurable_modes = tuple(
                mode for mode in value.modes if mode.damping_ratio is not None
            )
            mode_count = len(value.modes)
            composition = (
                "No EM mode detected" if not mode_count
                else "Single EM mode" if mode_count == 1
                else f"Multi-mode ({mode_count} EM modes)"
            )
            sheet.append((
                value.transfer_name,
                value.pre_fault_power,
                value.final_power,
                value.classification,
                mode_count,
                len(measurable_modes),
                composition,
                dominant.frequency_hz if dominant else None,
                dominant.damping_ratio * 100 if dominant and dominant.damping_ratio is not None else None,
                dominant.classification if dominant else None,
                critical.frequency_hz if critical else None,
                critical.damping_ratio * 100 if critical and critical.damping_ratio is not None else None,
                critical.classification if critical else None,
                sum(len(mode.damping_pairs) for mode in value.modes),
                self._damping_decision_basis(value.status),
                value.warning or value.comment,
            ))

    def _write_modes(self, sheet) -> None:
        sheet.append(("Transfer", "Candidate", "Frequency [Hz]", "Lower band [Hz]", "Upper band [Hz]", "Spectral power", "Spectral share [%]", "Representative", "Limiting", "Assessment", "Terminal direction", "Terminal damping [%]", "Selected damping [%]", "Selected extremum", "Envelope peak time [s]", "Envelope peak amplitude", "Envelope peak extremum", "Tail pairs", "Terminal sequence pairs", "Consecutive pairs", "Note"))
        for record in self.damping.records:
            value, dominant, critical = (
                record.assessment,
                record.assessment.dominant_mode,
                record.assessment.critical_mode,
            )
            total_spectral_power = sum(
                mode.spectral_power or 0.0 for mode in value.modes
            )
            for index, mode in enumerate(value.modes, start=1):
                spectral_share = (
                    mode.spectral_power / total_spectral_power * 100
                    if mode.spectral_power is not None and total_spectral_power
                    else None
                )
                sheet.append((value.transfer_name, index, mode.frequency_hz, mode.lower_band_hz, mode.upper_band_hz, mode.spectral_power, spectral_share, mode is dominant, mode is critical, mode.classification, mode.terminal_direction, mode.terminal_pair.damping_ratio * 100 if mode.terminal_pair else None, mode.damping_ratio * 100 if mode.damping_ratio is not None else None, mode.selected_extremum, mode.envelope_peak_time_s, mode.envelope_peak_amplitude, mode.envelope_peak_extremum, mode.tail_pair_count, mode.terminal_sequence_pair_count, len(mode.damping_pairs), mode.message))

    def _write_pairs(self, sheet) -> None:
        sheet.append(("Transfer", "Candidate", "Extremum", "Peak 1 time [s]", "Peak 2 time [s]", "A1", "A2", "Period [s]", "Damping [%]", "Peak 1 value", "Peak 2 value"))
        for record in self.damping.records:
            for candidate, mode in enumerate(record.assessment.modes, start=1):
                for pair in mode.damping_pairs:
                    sheet.append((record.assessment.transfer_name, candidate, pair.extremum, pair.first_time_s, pair.second_time_s, pair.first_amplitude, pair.second_amplitude, pair.period_s, pair.damping_ratio * 100, pair.first_value, pair.second_value))

    @staticmethod
    def _damping_decision_basis(status: str) -> str:
        """Explain the transfer classification without modifying the calculation."""
        return {
            "damped": "Every measurable EM mode meets the 5% damping limit.",
            "damped_under": "At least one measurable EM mode is below the 5% damping limit.",
            "growing": "At least one measurable EM mode has non-positive damping.",
            "not_measurable": "No detected EM mode had two usable consecutive extrema.",
            "inactive_pre_fault": "Transfer is inactive before the contingency.",
            "opened": "Transfer is open in the final operating condition.",
        }[status]

    def _write_high_frequency(self, sheet) -> None:
        sheet.append(("Transfer", "Component", "Frequency [Hz]", "Lower band [Hz]", "Upper band [Hz]", "Spectral power", "Status", "Note"))
        for record in self.high_frequency.records:
            assessment = record.assessment
            if assessment.components:
                for index, component in enumerate(assessment.components, start=1):
                    sheet.append((assessment.transfer_name, index, component.frequency_hz, component.lower_band_hz, component.upper_band_hz, component.spectral_power, "Warning", "Observation only"))
            else:
                sheet.append((assessment.transfer_name, None, None, None, None, None, "None detected", assessment.note))

    def _write_angular(self, sheet) -> None:
        sheet.append(("Machine", "Rotor-angle signal", "Inertia H [s]", "S [MVA]", "H·S [MVA*s]", "Minimum excursion [deg]", "Maximum excursion [deg]", "Maximum absolute excursion [deg]", "First ±120° violation [s]", "Within ±120°"))
        for machine in self.angular.machines:
            sheet.append((machine.machine_name, machine.signal_name, machine.inertia_s, machine.nominal_power_mva, machine.inertia_weight_mva_s, machine.minimum_excursion_degrees, machine.maximum_excursion_degrees, machine.maximum_absolute_excursion_degrees, machine.first_limit_violation_time_s, machine.is_within_limit))

    def _write_plant_recovery(self, sheet) -> None:
        """Write terminal P/Q recovery as a compact relational-style table."""
        sheet.append((
            "Plant",
            "Technology",
            "P signal",
            "Q signal",
            "P at 0.5 s [MW]",
            "P at tmax [MW]",
            "ΔP [MW]",
            "Max |ΔP| after evaluation [MW]",
            "Signed ΔP at max [MW]",
            "Max P-error time [s]",
            "P recovered (±5 MW)",
            "Q at 0.5 s [MVAr]",
            "Q at tmax [MVAr]",
            "ΔQ [MVAr]",
            "Max |ΔQ| after evaluation [MVAr]",
            "Signed ΔQ at max [MVAr]",
            "Max Q-error time [s]",
            "Q recovered (±5 MVAr)",
            "Evaluation starts [s]",
            "Final time [s]",
            "Terminal assessment",
        ))
        for record in self.plant_recovery.records:
            value = record.assessment
            sheet.append((
                record.plant_name,
                record.technology,
                record.active_power_signal,
                record.reactive_power_signal,
                value.pre_fault_active_power_mw,
                value.final_active_power_mw,
                value.active_power_error_mw,
                value.maximum_active_power_absolute_error_mw,
                value.maximum_active_power_error_mw,
                value.maximum_active_power_error_time_s,
                value.is_active_power_recovered,
                value.pre_fault_reactive_power_mvar,
                value.final_reactive_power_mvar,
                value.reactive_power_error_mvar,
                value.maximum_reactive_power_absolute_error_mvar,
                value.maximum_reactive_power_error_mvar,
                value.maximum_reactive_power_error_time_s,
                value.is_reactive_power_recovered,
                value.evaluation_start_time_s,
                value.final_time_s,
                value.status,
            ))

    def _write_spatial_stability(self, sheet) -> None:
        """Record the optional spatial-map outcome without relabelling stability."""
        sheet.append(("Item", "Value", "Detail"))
        sheet.append(("Map generation", self.spatial_stability.status, self.spatial_stability.reason))
        sheet.append(("Generated maps", len(self.spatial_stability.figure_paths), self.spatial_stability.output_directory))
        analysis = self.spatial_stability.analysis
        if analysis is not None and analysis.skipped_runs:
            sheet.append(("Run", "Spatial status", "Reason"))
            for item in analysis.skipped_runs:
                sheet.append((item.case_id, "Skipped", item.reason))

    def _write_definitions(self, sheet) -> None:
        voltage, damping = self.profile.voltage_criteria, self.profile.transfer_damping_criteria
        recovery = self.plant_recovery.criteria
        sheet.append(("Scope", "Parameter", "Value", "Unit / meaning"))
        rows = (
            ("Initialization", "Evaluation window", f"{self.initialization.criteria.evaluation_start_time_s:.3f}–{self.initialization.criteria.evaluation_end_time_s:.3f}", "s"),
            ("Initialization", "Smoothing duration", self.initialization.criteria.smoothing_duration_s, "s"),
            ("Initialization", "P99 fast-error limit", self.initialization.criteria.maximum_fast_relative_error_percentile_99 * 100, "% of voltage reference"),
            ("Voltage", "Crisis threshold", voltage.crisis_voltage_threshold_pu, "pu"),
            ("Voltage", "Recovery deadline after clearance", voltage.post_clearance_grace_period_s, "s"),
            ("Voltage", "Low-voltage threshold", voltage.low_voltage_threshold_pu, "pu"),
            ("Voltage", "Maximum cumulative time below 0.80 pu", voltage.maximum_cumulative_low_voltage_duration_s, "s"),
            ("Voltage", "Settlement peak-error limit", voltage.maximum_final_absolute_relative_error * 100, "%"),
            ("Voltage", "Settlement range limit", voltage.maximum_final_peak_to_peak_relative_range * 100, "%"),
            ("Damping", "EM frequency band", f"{damping.mode_minimum_frequency_hz:.1f}–{damping.mode_maximum_frequency_hz:.1f}", "Hz"),
            ("Damping", "Minimum damping", damping.minimum_damping_ratio * 100, "%"),
            ("Damping", "Representative EM mode", "Highest detected spectral power", "Descriptive transfer-level mode; it is not a weighted average of damping ratios."),
            ("Damping", "Limiting EM mode", "Lowest selected damping among measurable modes", "Traceability indicator for the conservative transfer assessment."),
            ("Damping", "Multi-mode transfer", "Two or more detected EM modes", "Modes remain independently measured; their damping ratios are not averaged."),
            ("Damping", "Spectral share", "Mode spectral power / total detected EM spectral power", "% within the detected EM modes; used only to identify the representative mode."),
            ("High frequency", "Observation band", f"{self.profile.high_frequency_criteria.minimum_frequency_hz:.1f}–{self.profile.high_frequency_criteria.maximum_frequency_hz:.1f}", "Hz; warning only"),
            ("Plant recovery", "Pre-fault reference", recovery.pre_fault_reference_time_s, "s; P and Q operating point"),
            ("Plant recovery", "Recovery observation starts", f"clearance + {recovery.recovery_evaluation_delay_s:.1f}", "s; all later samples remain visible"),
            ("Plant recovery", "Terminal active-power tolerance", recovery.active_power_tolerance_mw, "MW at tmax"),
            ("Plant recovery", "Terminal reactive-power tolerance", recovery.reactive_power_tolerance_mvar, "MVAr at tmax"),
            ("Plant recovery", "Terminal decision", "P and Q must both recover", "Observation only; does not alter dynamic stability"),
            ("Angular", "Centre of inertia", "Σ(δᵢHᵢSᵢ) / Σ(HᵢSᵢ)", "calculated per time step"),
            ("Angular", "Machine excursion", "δᵢ − δ_eje", "degrees"),
            ("Angular", "Stability criterion", "-120 < Δδᵢ(t) < +120", "strict inequality at every time step and for every machine"),
        )
        for row in rows:
            sheet.append(row)

    @staticmethod
    def _format_sheet(sheet) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, float):
                    cell.number_format = "0.000"

