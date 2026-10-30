"""Small case-index extension for the shared Monte Carlo Excel report."""

from pathlib import Path

from montecarlo.reports import AnalysisReport


def write_extreme_cases_report(
    output_directory: str | Path,
    *,
    study_analysis,
    transfer_analysis,
    high_frequency_analysis,
    dynamic_stability_analysis,
    angular_stability_analysis,
    scenario,
    profile,
    cases,
) -> Path:
    """Write shared analysis details plus deterministic-case identifiers."""
    report_path = AnalysisReport(
        study_analysis,
        transfer_analysis,
        high_frequency_analysis,
        dynamic_stability_analysis,
        angular_stability_analysis=angular_stability_analysis,
        scenario=scenario,
        profile=profile,
    ).write(output_directory)

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    workbook = load_workbook(report_path)
    sheet = workbook.create_sheet("Extreme Cases")
    sheet.append((
        "Run", "Case", "Common fn [Hz]", "Common xi [-]",
        "Voltage Stability (KIMAL_220)", "Transfer Damping", "Global Dynamic Stability",
    ))
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for run_id, case in enumerate(cases, start=1):
        dynamic = dynamic_stability_analysis.assessment_for(run_id)
        sheet.append((
            f"run_{run_id:06d}",
            case.definition.label,
            case.definition.fn_hz,
            case.definition.xi,
            study_analysis.status_for_run(run_id),
            transfer_analysis.modal_damping_status_for_run(run_id),
            dynamic.status,
        ))
    sheet.append(())
    if angular_stability_analysis is None:
        sheet.append((
            "Angular Stability", "Not evaluated", None, None, None, None,
            "Assets with inertia and nominal power were not provided.",
        ))
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[column[0].column_letter].width = min(width, 55)
    sheet.freeze_panes = "A2"
    workbook.save(report_path)
    workbook.close()
    return report_path
