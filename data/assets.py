"""Nominal-power lookup from the EMTP assets workbook."""

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Assets:
    """Provides nominal apparent power for EMTP assets."""

    workbook_path: Path
    _nominal_power_mva_by_name: dict[str, float] = field(init=False, repr=False)
    _inertia_s_by_name: dict[str, float] = field(init=False, repr=False)

    def __post_init__(self) -> None:
        self.workbook_path = Path(self.workbook_path).expanduser().resolve()

        if not self.workbook_path.is_file():
            raise FileNotFoundError(f"Assets workbook not found: {self.workbook_path}")

        (
            self._nominal_power_mva_by_name,
            self._inertia_s_by_name,
        ) = self._read_machine_properties()

    def get_nominal_power_mva(self, asset_name: str) -> float:
        """Return an asset's nominal apparent power in MVA."""
        try:
            return self._nominal_power_mva_by_name[asset_name]
        except KeyError as error:
            raise KeyError(f"Asset not found in the Assets workbook: {asset_name}") from error

    def get_inertia_s(self, asset_name: str) -> float:
        """Return the machine inertia constant H in seconds."""
        try:
            inertia_s = self._inertia_s_by_name[asset_name]
        except KeyError as error:
            raise KeyError(
                f"Asset has no inertia constant in the Assets workbook: {asset_name}"
            ) from error
        if inertia_s <= 0.0:
            raise ValueError(f"Asset {asset_name} has non-positive inertia: {inertia_s}")
        return inertia_s

    def get_inertia_weight_mva_s(self, asset_name: str) -> float:
        """Return the H·S weighting factor used in a centre-of-inertia angle."""
        return self.get_inertia_s(asset_name) * self.get_nominal_power_mva(asset_name)

    def _read_machine_properties(self) -> tuple[dict[str, float], dict[str, float]]:
        from openpyxl import load_workbook

        workbook = load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=True,
        )

        try:
            sheet = workbook["Unit Manager"]
        except KeyError as error:
            workbook.close()
            raise ValueError("Assets workbook does not contain a 'Unit Manager' sheet.") from error

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            workbook.close()
            raise ValueError("Assets workbook has no header row.")

        header_indexes = {header: index for index, header in enumerate(headers)}
        try:
            name_index = header_indexes["emtp_name"]
            power_index = header_indexes["App. Power [MVA]"]
            inertia_index = header_indexes["Inertia [s]"]
        except KeyError as error:
            workbook.close()
            raise ValueError(
                "Assets workbook must contain 'emtp_name', 'App. Power [MVA]', "
                "and 'Inertia [s]' columns."
            ) from error

        nominal_powers: dict[str, float] = {}
        inertias: dict[str, float] = {}
        try:
            for row in rows:
                asset_name = row[name_index]
                nominal_power_mva = row[power_index]
                inertia_s = row[inertia_index]

                if asset_name is None or nominal_power_mva is None:
                    continue

                name = str(asset_name).strip()
                if not name:
                    continue
                if name in nominal_powers:
                    raise ValueError(
                        f"Assets workbook contains a duplicate asset name: {name}"
                    )

                nominal_powers[name] = float(nominal_power_mva)
                if inertia_s is not None:
                    inertias[name] = float(inertia_s)
        finally:
            workbook.close()

        return nominal_powers, inertias
