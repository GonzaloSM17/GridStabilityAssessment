"""Excel reports produced from completed analysis results."""

from dataclasses import dataclass
from pathlib import Path

from montecarlo.batch_analysis import (
    ContingencyDynamicStabilityAnalysis,
    SignalStudyAnalysis,
    StudyAnalysis,
)
from stability.scenario import ContingencyScenario, EvaluationProfile
from stability.power_transfer import HighFrequencyStudyAnalysis, TransferStudyAnalysis


@dataclass(frozen=True)
class TransferDampingRunSummary:
    """One run's consolidated active-power transfer modal-damping conclusion."""

    status: str
    lowest_damping_percent: float | None
    lowest_transfer_name: str | None
    dominant_damping_percent: float | None
    dominant_transfer_name: str | None


@dataclass(frozen=True)
class AnalysisReport:
    """Write selected bus analyses and their consolidated run result."""

    study_analysis: StudyAnalysis
    transfer_analysis: TransferStudyAnalysis | None = None
    high_frequency_analysis: HighFrequencyStudyAnalysis | None = None
    dynamic_stability_analysis: ContingencyDynamicStabilityAnalysis | None = None
    scenario: ContingencyScenario | None = None
    profile: EvaluationProfile | None = None
    spatial_stability: object | None = None

    def write(self, output_directory: str | Path) -> Path:
        """Write the report at the root of the study analysis outputs."""
        from openpyxl import Workbook

        output_path = Path(output_directory) / "analysis_report.xlsx"
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        initialization_sheet = workbook.create_sheet("Initialization")
        stability_sheet = workbook.create_sheet("Stability")
        definitions_sheet = workbook.create_sheet("Definitions")
        if self.scenario is not None:
            scenario_sheet = workbook.create_sheet("Scenario")
        if self.transfer_analysis is not None:
            transfer_assessments_sheet = workbook.create_sheet("Transfer Assessments")
            transfer_modes_sheet = workbook.create_sheet("Electromechanical Modes")
            damping_pairs_sheet = workbook.create_sheet("Transfer Damping Pairs")
        if self.high_frequency_analysis is not None:
            high_frequency_sheet = workbook.create_sheet("High-Frequency Components")
        if self.spatial_stability is not None:
            spatial_sheet = workbook.create_sheet("Spatial Stability")

        self._write_summary(summary_sheet)
        self._write_initialization(initialization_sheet)
        self._write_stability(stability_sheet)
        self._write_definitions(definitions_sheet)
        if self.scenario is not None:
            self._write_scenario(scenario_sheet)
        if self.transfer_analysis is not None:
            self._write_transfer_assessments(transfer_assessments_sheet)
            self._write_transfer_modes(transfer_modes_sheet)
            self._write_transfer_damping_pairs(damping_pairs_sheet)
        if self.high_frequency_analysis is not None:
            self._write_high_frequency_components(high_frequency_sheet)
        if self.spatial_stability is not None:
            self._write_spatial_stability(spatial_sheet)
        workbook.save(output_path)
        workbook.close()
        return output_path

    def _write_scenario(self, sheet) -> None:
        """Document the physical contingency and operating state under evaluation."""
        sheet.append(("Field", "Value"))
        values = (
            ("Name", self.scenario.name),
            ("Severity", self.scenario.severity),
            ("Fault type", self.scenario.fault_type),
            ("Fault time [s]", self.scenario.fault_time_s),
            ("Clearance time [s]", self.scenario.clearance_time_s),
            ("Location", self.scenario.location or "-"),
            ("Operating state", self.scenario.operating_state),
        )
        for value in values:
            sheet.append(value)
        self._format_sheet(sheet)

    def _write_summary(self, sheet) -> None:
        """Write a compact run index; detailed records live in the data sheets."""
        sheet.append((
            "Run",
            "Contingency Dynamic Stability",
            "Dynamic Stability Reason",
            "Voltage Stability",
            "Active-Power Transfer Modal Damping",
            "Worst Dominant EM Damping [%]",
            "Dominant EM Transfer",
            "Lowest EM Mode Damping [%]",
            "Lowest EM Mode Transfer",
            "High-Frequency Components",
            "Spatial Map",
            "Spatial Map Note",
        ))
        for run_id in self.study_analysis.run_ids:
            damping = self._transfer_damping_summary_for_run(run_id)
            dynamic_stability = (
                self.dynamic_stability_analysis.assessment_for(run_id)
                if self.dynamic_stability_analysis is not None
                else None
            )
            high_frequency_count = self._high_frequency_component_count(run_id)
            spatial_status, spatial_note = self._spatial_status_for_run(run_id)
            sheet.append((
                self._run_name(run_id),
                dynamic_stability.status if dynamic_stability is not None else "Not evaluated",
                dynamic_stability.reason if dynamic_stability is not None else None,
                self.study_analysis.status_for_run(run_id),
                damping.status,
                damping.dominant_damping_percent,
                damping.dominant_transfer_name,
                damping.lowest_damping_percent,
                damping.lowest_transfer_name,
                f"Detected ({high_frequency_count})" if high_frequency_count else "None detected",
                spatial_status,
                spatial_note,
            ))
        self._format_sheet(sheet)

    def _write_transfer_assessments(self, sheet) -> None:
        """Write one database-style assessment row per run and transfer."""
        sheet.append((
            "Run", "Active-Power Transfer", "Pre-fault P", "Final P", "Assessment",
            "Dominant EM Mode [Hz]", "Dominant EM Damping [%]",
            "Lowest EM Mode [Hz]", "Lowest EM Mode Damping [%]",
            "High-Frequency Components [Hz]", "Note",
        ))
        for item in self.transfer_analysis.assessments:
            mode = item.critical_mode
            dominant_mode = item.dominant_mode
            high_frequency = self._high_frequency_assessment(
                item.run_id,
                item.transfer_name,
            )
            sheet.append((
                self._run_name(item.run_id), item.transfer_name,
                item.pre_fault_power, item.final_power, item.classification,
                dominant_mode.frequency_hz if dominant_mode else None,
                (
                    dominant_mode.damping_ratio * 100
                    if dominant_mode and dominant_mode.damping_ratio is not None
                    else None
                ),
                mode.frequency_hz if mode else None,
                mode.damping_ratio * 100 if mode and mode.damping_ratio is not None else None,
                self._component_frequency_list(high_frequency),
                item.warning or item.comment,
            ))
        self._format_transfer_sheet(sheet)

    def _write_high_frequency_components(self, sheet) -> None:
        """Write observation-only spectral components outside the EM band."""
        sheet.append((
            "Run", "Transfer", "Component", "Frequency [Hz]",
            "Lower Band [Hz]", "Upper Band [Hz]", "Spectral Power", "Status", "Note",
        ))
        for assessment in self.high_frequency_analysis.assessments:
            if assessment.components:
                for index, component in enumerate(assessment.components, start=1):
                    sheet.append((
                        self._run_name(assessment.run_id),
                        assessment.transfer_name,
                        index,
                        component.frequency_hz,
                        component.lower_band_hz,
                        component.upper_band_hz,
                        component.spectral_power,
                        "Detected",
                        None,
                    ))
            else:
                sheet.append((
                    self._run_name(assessment.run_id),
                    assessment.transfer_name,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "None detected",
                    assessment.note,
                ))
        self._format_transfer_sheet(sheet)

    def _write_transfer_modes(self, sheet) -> None:
        """Write one database-style row per electromechanical mode."""
        sheet.append((
            "Run", "Transfer", "Candidate", "Frequency [Hz]", "Lower Band [Hz]", "Upper Band [Hz]",
            "Spectral Power", "Dominant", "Assessment", "Terminal Direction",
            "Terminal Damping [%]", "Selected Damping [%]", "Selected Extremum",
            "Envelope Peak Time [s]", "Envelope Peak Amplitude", "Envelope Peak Extremum",
            "Tail Pairs", "Terminal Sequence Pairs", "Consecutive Pairs", "Note",
        ))
        for item in self.transfer_analysis.assessments:
            dominant_mode = item.dominant_mode
            for index, mode in enumerate(item.modes, start=1):
                sheet.append((
                    self._run_name(item.run_id), item.transfer_name, index, mode.frequency_hz,
                    mode.lower_band_hz, mode.upper_band_hz,
                    mode.spectral_power,
                    mode is dominant_mode,
                    mode.classification,
                    mode.terminal_direction,
                    (
                        mode.terminal_pair.damping_ratio * 100
                        if mode.terminal_pair is not None
                        else None
                    ),
                    mode.damping_ratio * 100 if mode.damping_ratio is not None else None,
                    mode.selected_extremum,
                    mode.envelope_peak_time_s, mode.envelope_peak_amplitude,
                    mode.envelope_peak_extremum, mode.tail_pair_count,
                    mode.terminal_sequence_pair_count,
                    len(mode.damping_pairs), mode.message,
                ))
        self._format_transfer_sheet(sheet)

    def _write_transfer_damping_pairs(self, sheet) -> None:
        """Write every consecutive same-kind extrema pair used by each mode."""
        sheet.append((
            "Run", "Transfer", "Candidate", "Frequency [Hz]", "Selected", "Extremum",
            "Peak 1 Time [s]", "Peak 2 Time [s]", "Period [s]",
            "Value 1", "Value 2", "A1", "A2", "A2 / A1", "Damping [%]",
        ))
        for item in self.transfer_analysis.assessments:
            for mode_index, mode in enumerate(item.modes, start=1):
                for pair in mode.damping_pairs:
                    sheet.append((
                        self._run_name(item.run_id), item.transfer_name, mode_index,
                        mode.frequency_hz,
                        mode.status != "not_measurable" and pair == mode.selected_pair,
                        pair.extremum,
                        pair.first_time_s, pair.second_time_s, pair.period_s,
                        pair.first_value, pair.second_value,
                        pair.first_amplitude, pair.second_amplitude, pair.amplitude_ratio,
                        pair.damping_ratio * 100,
                    ))
        self._format_transfer_sheet(sheet)

    def _transfer_damping_summary_for_run(self, run_id: int) -> TransferDampingRunSummary:
        """Return the limiting active-power transfer modal-damping state."""
        if self.transfer_analysis is None:
            return TransferDampingRunSummary("Not measurable", None, None, None, None)
        assessments = [
            item for item in self.transfer_analysis.assessments
            if item.run_id == run_id and item.status not in {"opened", "inactive_pre_fault"}
        ]
        damping = [item.damping_ratio for item in assessments if item.damping_ratio is not None]
        if not assessments:
            return TransferDampingRunSummary("Not applicable", None, None, None, None)
        status = self.transfer_analysis.modal_damping_status_for_run(run_id)
        lowest_assessment = min(
            (item for item in assessments if item.damping_ratio is not None),
            key=lambda item: item.damping_ratio,
            default=None,
        )
        dominant_assessment = min(
            (
                item
                for item in assessments
                if item.dominant_damping_ratio is not None
            ),
            key=lambda item: item.dominant_damping_ratio,
            default=None,
        )
        return TransferDampingRunSummary(
            status=status,
            lowest_damping_percent=(
                lowest_assessment.damping_ratio * 100
                if lowest_assessment is not None
                else None
            ),
            lowest_transfer_name=(
                lowest_assessment.transfer_name if lowest_assessment is not None else None
            ),
            dominant_damping_percent=(
                dominant_assessment.dominant_damping_ratio * 100
                if dominant_assessment is not None
                else None
            ),
            dominant_transfer_name=(
                dominant_assessment.transfer_name if dominant_assessment is not None else None
            ),
        )


    def _high_frequency_assessment(self, run_id: int, transfer_name: str):
        if self.high_frequency_analysis is None:
            return None
        return self.high_frequency_analysis.assessment_for(run_id, transfer_name)

    def _high_frequency_component_count(self, run_id: int) -> int:
        if self.high_frequency_analysis is None:
            return 0
        return sum(
            len(assessment.components)
            for assessment in self.high_frequency_analysis.assessments
            if assessment.run_id == run_id
        )

    @staticmethod
    def _component_frequency_list(assessment) -> str | None:
        if assessment is None or not assessment.components:
            return None
        return "; ".join(
            f"{component.frequency_hz:.2f}"
            for component in assessment.components
        )

    def _write_spatial_stability(self, sheet) -> None:
        """Document optional map generation and any skipped Monte Carlo runs."""
        result = self.spatial_stability
        sheet.append(("Item", "Value", "Detail"))
        sheet.append(("Map generation", result.status, result.reason))
        sheet.append(("Generated maps", len(result.figure_paths), str(result.output_directory)))
        analysis = result.analysis
        if analysis is not None:
            sheet.append(("Mapped controller points", len(analysis.points), None))
            sheet.append(("Excluded controller devices", len(analysis.excluded_devices), None))
            sheet.append(("Skipped runs", len(analysis.skipped_runs), None))
            if analysis.skipped_runs:
                sheet.append(("Run", "Spatial status", "Reason"))
                for item in analysis.skipped_runs:
                    sheet.append((item.case_id, "Skipped", item.reason))

    def _spatial_status_for_run(self, run_id: int) -> tuple[str, str | None]:
        if self.spatial_stability is None:
            return "Not generated", None
        result = self.spatial_stability
        run_name = self._run_name(run_id)
        analysis = result.analysis
        if analysis is None:
            return result.status, result.reason
        skipped = next(
            (item for item in analysis.skipped_runs if item.case_id == run_name),
            None,
        )
        if skipped is not None:
            return "Skipped", skipped.reason
        mapped = any(point.case_id == run_name for point in analysis.points)
        if mapped:
            return "Mapped", None
        return "Skipped", result.reason or "No spatial controller points were available."

    def _write_initialization(self, sheet) -> None:
        sheet.append(
            (
                "Bus",
                "Signal",
                "Run",
                "Status",
                "P99 Fast Error [%]",
                "Limit [%]",
                "Reason",
                "Reference Voltage [pu]",
            )
        )
        for bus_analysis in self.study_analysis.bus_analyses:
            voltage_analysis = bus_analysis.signal_analysis
            criteria = self._initialization_criteria(voltage_analysis)
            for result in voltage_analysis.run_results:
                initialization = result.initialization
                sheet.append(
                    (
                        bus_analysis.bus_name,
                        voltage_analysis.signal_name,
                        self._run_name(result.run_id),
                        (
                            "Well initialized"
                            if initialization.is_well_initialized
                            else "Poorly initialized"
                        ),
                        initialization.fast_relative_error_percentile_99 * 100,
                        criteria.maximum_fast_relative_error_percentile_99 * 100,
                        ", ".join(initialization.rejection_reasons) or "-",
                        initialization.reference_value,
                    )
                )

        self._format_sheet(sheet)

    def _write_stability(self, sheet) -> None:
        sheet.append(
            (
                "Bus",
                "Signal",
                "Run",
                "Status",
                "Failed Criteria",
                "0.70 pu Decision Time [s]",
                "0.70 pu Decision After Clearance [s]",
                "0.70 pu Decision Event",
                "0.70 pu Deadline After Clearance [s]",
                "0.70 pu Status",
                "Cumulative V < 0.80 pu [s]",
                "0.80 pu Duration Limit [s]",
                "0.80 pu Status",
                "Final Settlement Status",
                "Fault-on Area [pu*s]",
                "Recovery Area [pu*s]",
                "Unified Response Area [pu*s]",
                "Final Error [%]",
                "Final Error Limit [%]",
                "Final Range [%]",
                "Final Range Limit [%]",
            )
        )
        for bus_analysis in self.study_analysis.bus_analyses:
            voltage_analysis = bus_analysis.signal_analysis
            for result in voltage_analysis.run_results:
                stability = result.stability
                if stability is None:
                    reason = (
                        "Initialization failed"
                        if not result.initialization.is_well_initialized
                        else "Simulation duration is insufficient"
                    )
                    sheet.append(
                        (
                            bus_analysis.bus_name,
                            voltage_analysis.signal_name,
                            self._run_name(result.run_id),
                            "Not evaluated",
                            reason,
                        )
                    )
                    continue

                is_stable = stability.is_stable
                sheet.append(
                    (
                        bus_analysis.bus_name,
                        voltage_analysis.signal_name,
                        self._run_name(result.run_id),
                        "Stable" if is_stable else "Unstable",
                        ", ".join(stability.failed_criteria) if not is_stable else "-",
                        stability.crisis_decision_time_s,
                        stability.crisis_decision_duration_s,
                        stability.crisis_decision_event,
                        stability.criteria.post_clearance_grace_period_s,
                        self._pass_fail(stability.is_crisis_voltage_compliant),
                        stability.cumulative_low_voltage_duration_s,
                        stability.criteria.maximum_cumulative_low_voltage_duration_s,
                        self._pass_fail(stability.is_low_voltage_duration_compliant),
                        self._pass_fail(stability.is_settled),
                        stability.fault_on_area_pu_s,
                        stability.recovery_area_pu_s,
                        stability.unified_response_area_pu_s,
                        stability.final_settlement_peak_absolute_relative_error * 100,
                        stability.criteria.maximum_final_absolute_relative_error * 100,
                        stability.final_settlement_peak_to_peak_relative_range * 100,
                        stability.criteria.maximum_final_peak_to_peak_relative_range * 100,
                    )
                )

        self._format_sheet(sheet)

    def _write_definitions(self, sheet) -> None:
        reference_analysis = self.study_analysis.bus_analyses[0].signal_analysis
        initialization = self._initialization_criteria(reference_analysis)
        stability = self._stability_criteria(reference_analysis)
        sheet.append(("Area", "Definition / Criterion", "Value"))
        rows = [
            ("Initialization", "Evaluation window [s]", f"{initialization.evaluation_start_time_s:.2f} to {initialization.evaluation_end_time_s:.2f}"),
            ("Initialization", "Smoothing window [s]", initialization.smoothing_duration_s),
            ("Initialization", "Maximum P99 fast error [%]", initialization.maximum_fast_relative_error_percentile_99 * 100),
            ("Stability", "Minimum simulation end time [s]", stability.minimum_simulation_end_time_s),
            ("Voltage", "Fault time [s]", stability.fault_time_s),
            ("Voltage", "Clearance time [s]", stability.clearance_time_s),
            ("Voltage", "Response end time [s]", stability.response_end_time_s),
            ("Voltage", "0.70 pu crisis threshold [pu]", stability.crisis_voltage_threshold_pu),
            ("Voltage", "0.70 pu deadline after clearance [s]", stability.post_clearance_grace_period_s),
            ("Voltage", "Low-voltage threshold [pu]", stability.low_voltage_threshold_pu),
            ("Voltage", "Maximum cumulative duration below 0.80 pu [s]", stability.maximum_cumulative_low_voltage_duration_s),
            ("Stability", "Final absolute error limit [%]", stability.maximum_final_absolute_relative_error * 100),
            ("Stability", "Final range limit [%]", stability.maximum_final_peak_to_peak_relative_range * 100),
        ]
        if self.transfer_analysis and self.transfer_analysis.assessments:
            transfer_criteria = self.transfer_analysis.assessments[0].criteria
            rows.extend(
                (
                    (
                        "Active-Power Transfer Modal Damping",
                        "Electromechanical response band [Hz]",
                        f"{transfer_criteria.mode_minimum_frequency_hz:.1f} to "
                        f"{transfer_criteria.mode_maximum_frequency_hz:.1f}",
                    ),
                    (
                        "Active-Power Transfer Modal Damping",
                        "Post-clearance guard time [s]",
                        transfer_criteria.post_clearance_guard_time_s,
                    ),
                    (
                        "Active-Power Transfer Modal Damping",
                        "Analysis start time [s]",
                        transfer_criteria.analysis_start_time_s,
                    ),
                    (
                        "Active-Power Transfer Modal Damping",
                        "Mode detection",
                        "Every local spectrum maximum in the response band; the global maximum is included at a band edge.",
                    ),
                    (
                        "Active-Power Transfer Modal Damping",
                        "Mode bands",
                        "Each candidate is bounded by its nearest spectral valleys; outer candidates use the response-band limits when a valley is unavailable.",
                    ),
                    ("Active-Power Transfer Modal Damping", "Reference", "P_final is the median transfer power in the final window."),
                    ("Active-Power Transfer Modal Damping", "Measured signal", "P(t) - P_final; each spectral mode is evaluated independently."),
                    ("Active-Power Transfer Modal Damping", "Dominant EM mode", "The detected EM mode with the greatest spectral power."),
                    ("Active-Power Transfer Modal Damping", "Lowest EM mode", "The detected EM mode with the smallest selected damping ratio."),
                    ("Active-Power Transfer Modal Damping", "Damping pairs", "Consecutive local maxima or consecutive local minima of the isolated EM signal, following the NT logarithmic-decrement method."),
                    ("Active-Power Transfer Modal Damping", "Extremum amplitude", "Each local extremum keeps its signed signal value and uses its absolute distance from P_final for the logarithmic decrement. Zero crossings are retained."),
                    ("Active-Power Transfer Modal Damping", "Damping phase", "Contains consecutive pairs of the same local-extremum kind after the maximum observed EM amplitude."),
                    ("Active-Power Transfer Modal Damping", "Terminal sequence", "Starts at the final pair and extends backwards only while its envelope direction is unchanged."),
                    ("Active-Power Transfer Modal Damping", "Selected damping", "Decreasing terminal sequence: largest positive xi. Non-decreasing terminal sequence: most negative xi."),
                    ("Active-Power Transfer Modal Damping", "Direction boundary", "An earlier pair is not selected across an intervening pair with the opposite envelope direction."),
                    ("Active-Power Transfer Modal Damping", "Available evidence", "One valid pair in the terminal sequence is sufficient; both tail and terminal-sequence counts are reported."),
                    ("Active-Power Transfer Modal Damping", "Damped", f"All measurable modes have xi >= {transfer_criteria.minimum_damping_ratio * 100:.0f}%"),
                    ("Active-Power Transfer Modal Damping", "Damped under", f"At least one measurable mode has 0% < xi < {transfer_criteria.minimum_damping_ratio * 100:.0f}%"),
                    ("Active-Power Transfer Modal Damping", "Growing", "At least one measurable mode ends with a non-decreasing terminal sequence (xi <= 0%)."),
                    ("Active-Power Transfer Modal Damping", "Not measurable", "No detected mode has a valid consecutive local-extremum pair after its envelope maximum."),
                    ("Active-Power Transfer Modal Damping", "Minimum damping ratio [%]", transfer_criteria.minimum_damping_ratio * 100),
                )
            )
        if self.high_frequency_analysis and self.high_frequency_analysis.assessments:
            high_frequency_criteria = self.high_frequency_analysis.assessments[0].criteria
            rows.extend(
                (
                    (
                        "High-frequency components",
                        "Observation band [Hz]",
                        f"{high_frequency_criteria.minimum_frequency_hz:.1f} to "
                        f"{high_frequency_criteria.maximum_frequency_hz:.1f}",
                    ),
                    (
                        "High-frequency components",
                        "Measured signal",
                        "P(t) - P_final; detected independently from EM damping.",
                    ),
                    (
                        "High-frequency components",
                        "Detection",
                        "Local spectrum maxima bounded by their nearest spectral valleys.",
                    ),
                    (
                        "High-frequency components",
                        "Assessment",
                        "Observation only: no damping ratio or stability classification is assigned.",
                    ),
                )
            )
        rows.extend(
            (
                (
                    "Contingency Dynamic Stability",
                    "Scope",
                    "Combined voltage stability and active-power transfer modal damping for the defined contingency.",
                ),
                (
                    "Contingency Dynamic Stability",
                    "Stable",
                    "Voltage stability is Stable and active-power transfer modal damping is Damped.",
                ),
                (
                    "Contingency Dynamic Stability",
                    "Unstable",
                    "Either voltage stability is Unstable or active-power transfer modal damping is Growing or Damped under.",
                ),
                (
                    "Contingency Dynamic Stability",
                    "Not evaluated",
                    "Either voltage stability or active-power transfer modal damping was not evaluated.",
                ),
            )
        )
        for row in rows:
            sheet.append(row)

        self._format_sheet(sheet)

    def _initialization_criteria(self, voltage_analysis: SignalStudyAnalysis):
        if not voltage_analysis.run_results:
            if self.profile is not None:
                return self.profile.initialization_criteria
            from stability.initialization import InitializationCriteria

            return InitializationCriteria()
        return voltage_analysis.run_results[0].initialization.criteria

    def _stability_criteria(self, voltage_analysis: SignalStudyAnalysis):
        for result in voltage_analysis.run_results:
            if result.stability is not None:
                return result.stability.criteria
        if self.profile is not None:
            return self.profile.voltage_criteria
        from stability.voltage import StabilityCriteria

        return StabilityCriteria()

    @staticmethod
    def _run_name(run_id: int) -> str:
        return f"run_{run_id:05d}"

    @staticmethod
    def _pass_fail(value: bool) -> str:
        return "Pass" if value else "Fail"

    @staticmethod
    def _format_sheet(sheet) -> None:
        from openpyxl.styles import Font

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.###"
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 45)

    @staticmethod
    def _format_transfer_sheet(sheet) -> None:
        AnalysisReport._format_sheet(sheet)
        for column in sheet.columns:
            for cell in column[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
