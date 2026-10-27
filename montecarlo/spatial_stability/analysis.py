"""Adapt published run parameters and shared stability analyses to spatial maps."""
from collections import Counter
from dataclasses import replace
from math import isfinite, pi, sqrt
from pathlib import Path
import re
import unicodedata

from data.assets import Assets
from data.run import MonteCarloRun
from stability.dynamic_stability import assess_dynamic_stability
from stability.scenario import EvaluationProfile
from .models import (
    SpatialControllerPoint,
    SpatialCaseAssessment,
    SpatialSkippedRun,
    SpatialStabilityAnalysis,
    SpatialStabilityResult,
)

# Explicit rename confirmed by the user; exact workbook names take precedence.
ASSET_ALIASES = {"PFV_POZOALMONTE_HUAYCA": "PFV_LA_HUAYCA_II"}


def normalized_name(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]", "", value.upper())


class SpatialAssetLookup:
    """Read drawing coordinates without altering the existing Assets API."""
    def __init__(self, assets: Assets, sheet="Unit Manager", x_column="posX", y_column="posY"):
        from openpyxl import load_workbook
        self.assets = assets
        self.positions = {}
        self.names = {}
        workbook = load_workbook(assets.workbook_path, read_only=True, data_only=True)
        try:
            if sheet not in workbook.sheetnames:
                raise ValueError(f"Assets workbook has no sheet {sheet!r}")
            rows = workbook[sheet].iter_rows(values_only=True)
            headers = next(rows, ())
            required = ("emtp_name", x_column, y_column)
            missing = [name for name in required if name not in headers]
            if missing:
                raise ValueError(f"Spatial maps require workbook columns: {missing}")
            ni, xi, yi = [headers.index(name) for name in required]
            for row in rows:
                if row[ni] is None or not str(row[ni]).strip(): continue
                name = str(row[ni]).strip()
                key = normalized_name(name)
                if key in self.names:
                    raise ValueError(f"Ambiguous normalized asset names: {name!r} and {self.names[key]!r}")
                self.names[key] = name
                try:
                    x, y = float(row[xi]), float(row[yi])
                except (TypeError, ValueError):
                    continue
                if isfinite(x) and isfinite(y): self.positions[name] = (x, y)
        finally:
            workbook.close()

    def resolve(self, generator_name):
        key = normalized_name(generator_name)
        name = self.names.get(key)
        if name is None:
            base = re.sub(r"\[\d+\]$", "", generator_name)
            name = self.names.get(normalized_name(base))
            if name is None:
                alias = ASSET_ALIASES.get(base.upper())
                if alias: name = self.names.get(normalized_name(alias))
        if name is None: raise KeyError(f"No asset matches {generator_name!r}")
        if name not in self.positions: raise ValueError(f"Asset {name!r} has no finite posX/posY")
        power = self.assets.get_nominal_power_mva(name)
        if not isfinite(power) or power <= 0: raise ValueError(f"Asset {name!r} has invalid MVA")
        return name, *self.positions[name], power


def read_controller_points(run, lookup):
    """Use active GeneratorName/Kp/Ki rows from the native semicolon CSV.

    Repeated plants require explicit UnitId values; their plant MVA is divided
    between active controller rows to avoid multiplying the plant's capacity.
    """
    parameters = run.read_parameters()
    missing = {"GeneratorName", "InService", "Kp", "Ki"}.difference(parameters.columns)
    if missing: raise ValueError(f"{run.display_name}: parameters.csv missing {sorted(missing)}")
    active, excluded = [], []
    for row in parameters.iter_rows(named=True):
        if row["InService"] != 1: continue
        if row["Kp"] is None or row["Ki"] is None: continue  # Same controller selection as data.run.
        if row["GeneratorName"] is None or not str(row["GeneratorName"]).strip():
            raise ValueError(f"{run.display_name}: active controller without GeneratorName")
        generator = str(row["GeneratorName"]).strip()
        kp, ki = float(row["Kp"]), float(row["Ki"])
        if not isfinite(kp) or not isfinite(ki) or ki <= 0:
            raise ValueError(f"{run.display_name}: nonfinite Kp or non-positive Ki for {generator}")
        try:
            asset, x, y, power = lookup.resolve(generator)
        except (KeyError, ValueError) as error:
            excluded.append(dict(case_id=run.display_name, generator_name=generator, reason=str(error)))
            continue
        unit = row.get("UnitId")
        if unit is None or not str(unit).strip(): unit = row.get("UnitID")
        device_id = f"{generator}::UnitId={unit}" if unit is not None and str(unit).strip() else generator
        active.append((device_id, generator, asset, x, y, power, kp, ki))
    keys = [row[0] for row in active]
    if len(keys) != len(set(keys)):
        raise ValueError(f"{run.display_name}: duplicate controller identity; provide distinct UnitId values")
    counts = Counter(row[2] for row in active)
    points = tuple(SpatialControllerPoint(
        case_id=run.display_name, run_id=run.run_id, device_id=identity,
        generator_name=generator, asset_name=asset, x=x, y=y,
        nominal_power_mva=power/counts[asset], kp=kp, ki=ki,
        xi=kp/(2*sqrt(ki)), fn_hz=sqrt(ki)/(2*pi),
    ) for identity, generator, asset, x, y, power, kp, ki in active)
    if not points:
        reasons = "; ".join(item["reason"] for item in excluded)
        raise ValueError(f"{run.display_name}: no active controllers with usable coordinates. {reasons}")
    return points, tuple(excluded)


def assess_components(run_id, voltage, damping, angular):
    """Apply the same aggregation rules as SingleRunWorkbench.create_report."""
    voltage_status = "Not evaluated"
    if voltage is not None and voltage.assessments:
        voltage_status = "Stable" if all(item.assessment.is_stable for item in voltage.assessments) else "Unstable"
    damping_status = "Not measurable"
    if damping is not None:
        usable = [record.assessment for record in damping.records
                  if record.assessment.status not in {"opened", "inactive_pre_fault"}]
        if not usable: damping_status = "Not applicable"
        elif any(item.status == "growing" for item in usable): damping_status = "Growing"
        elif any(item.status == "damped_under" for item in usable): damping_status = "Damped under"
        elif any(item.damping_ratio is not None for item in usable): damping_status = "Damped"
    angular_status = "Not evaluated"
    if angular is not None and angular.machines:
        angular_status = "Stable" if angular.is_stable else "Unstable"
    return assess_dynamic_stability(run_id, voltage_status, damping_status, angular_status)


def assess_run(run, assets, profile, *, voltage=None, damping=None, angular=None):
    """Reuse supplied analyses or calculate each using the new repository's code."""
    from single_run.voltage_stability import SingleRunVoltageStability
    from single_run.damping_stability import SingleRunDampingStability
    from single_run.angular_stability import SingleRunAngularStability
    diagnostics = []
    analyses = dict(voltage=voltage, damping=damping, angular=angular)
    calculators = dict(
        voltage=lambda: SingleRunVoltageStability(run, profile).analyze(),
        damping=lambda: SingleRunDampingStability(run, profile).analyze(),
        angular=lambda: SingleRunAngularStability(run, profile.scenario, assets).analyze(),
    )
    for name, calculate in calculators.items():
        if analyses[name] is not None: continue
        if run.dataset is None:
            diagnostics.append(f"{name}: run has no Parquet dataset")
            continue
        try:
            analyses[name] = calculate()
        except (ValueError, KeyError, FileNotFoundError) as error:
            diagnostics.append(f"{name}: {error}")
    assessment = assess_components(run.run_id, **analyses)
    return SpatialCaseAssessment(run.display_name, assessment, tuple(diagnostics))


class SpatialStabilityAnalyzer:
    def __init__(self, assets: Assets, profile: EvaluationProfile):
        self.assets, self.profile = assets, profile
        self.lookup = SpatialAssetLookup(assets)

    def analyze(self, runs, *, component_analyses=None):
        runs = tuple(runs)
        if not runs: raise ValueError("No runs were selected")
        case_ids = [run.display_name for run in runs]
        if len(case_ids) != len(set(case_ids)): raise ValueError("Run folder names must be unique")
        points, cases, excluded, skipped_runs = [], [], [], []
        component_analyses = component_analyses or {}
        for run in runs:
            try:
                run_points, omitted = read_controller_points(run, self.lookup)
            except (FileNotFoundError, KeyError, ValueError) as error:
                skipped_runs.append(SpatialSkippedRun(run.display_name, str(error)))
                run_points, omitted = (), ()
            points.extend(run_points); excluded.extend(omitted)
            cases.append(assess_run(run, self.assets, self.profile, **component_analyses.get(run.display_name, {})))
        return SpatialStabilityAnalysis(
            tuple(points),
            tuple(cases),
            tuple(excluded),
            tuple(skipped_runs),
        )

    def create(self, runs, output_directory, *, component_analyses=None):
        from .plotting import export_maps
        runs = tuple(runs)
        analysis = self.analyze(runs, component_analyses=component_analyses)
        if not analysis.points:
            reason = (
                analysis.skipped_runs[0].reason
                if len(analysis.skipped_runs) == 1
                else "No run has active controllers with usable spatial metadata."
            )
            return self._skipped_result(output_directory, reason, analysis)
        result = export_maps(analysis, output_directory, self.assets, self.profile, runs)
        status = (
            "Completed with warnings"
            if analysis.skipped_runs or analysis.excluded_devices
            else "Completed"
        )
        reason = (
            "Some runs or controller devices could not be mapped."
            if status == "Completed with warnings"
            else None
        )
        return replace(result, status=status, reason=reason)

    @classmethod
    def create_optional(cls, assets, profile, runs, output_directory, *, component_analyses=None):
        """Create spatial artefacts without allowing metadata issues to stop a flow."""
        runs = tuple(runs)
        try:
            return cls(assets, profile).create(
                runs,
                output_directory,
                component_analyses=component_analyses,
            )
        except (FileNotFoundError, KeyError, OSError, ValueError) as error:
            analysis = SpatialStabilityAnalysis(
                (),
                (),
                (),
                tuple(SpatialSkippedRun(run.display_name, str(error)) for run in runs),
            )
            return cls._skipped_result(output_directory, str(error), analysis)

    @staticmethod
    def _skipped_result(output_directory, reason, analysis):
        output = Path(output_directory).expanduser().resolve()
        return SpatialStabilityResult(
            output,
            analysis,
            (),
            status="Skipped",
            reason=reason,
        )
